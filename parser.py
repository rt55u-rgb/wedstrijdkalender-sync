from openpyxl import load_workbook
from datetime import datetime

from config import EXCEL_FILE
from models import Match


def _clean(value):
    if value is None:
        return ""

    if isinstance(value, str):
        return value.strip()

    return str(value)


def load_matches():
    wb = load_workbook(EXCEL_FILE, data_only=True)
    ws = wb["Blad1"]

    matches = []

    current = None

    for row in ws.iter_rows(min_row=2, values_only=True):

        day = row[0]
        date = row[1]
        category = _clean(row[2])
        system = _clean(row[3])
        info = _clean(row[4])
        location = _clean(row[5])
        start = row[6]
        age = _clean(row[7])
        notes = _clean(row[8])

        # Nieuwe wedstrijd
        if date:

            if current:
                matches.append(current)

            title = category

            if info:
                title += f" - {info}"

            uid = (
                date.strftime("%Y%m%d")
                + "-"
                + category
                + "-"
                + info
            ).replace(" ", "_")

            current = Match(
                uid=uid,
                date=date,
                title=title,
                category=category,
                system=system,
                info=info,
                location=location,
                start_time=start.strftime("%H:%M") if start else "",
                age=age,
                notes=notes,
            )

        else:

            if current is None:
                continue

            # Locatie loopt soms door
            if location:
                current.location += "\n" + location

            # Overige info loopt vaak door
            if notes:
                current.notes += "\n" + notes

    if current:
        matches.append(current)

    return matches


if __name__ == "__main__":

    wedstrijden = load_matches()

    print(f"{len(wedstrijden)} wedstrijden gevonden\n")

    for w in wedstrijden[:5]:
        print("=" * 60)
        print(w.title)
        print(w.date)
        print(w.location)
        print(w.notes)
